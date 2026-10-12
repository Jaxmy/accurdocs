"""DOCX 解析：段落流 + 表格。"""

from __future__ import annotations

from dataclasses import dataclass, field

from ...models.document import TableBlock

__all__ = ["DocxParseResult", "parse_docx"]


@dataclass
class DocxParseResult:
    pages: list[tuple[str, int | None]] = field(default_factory=list)
    tables: list[TableBlock] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    ocr_required: bool = False


def parse_docx(path: str) -> DocxParseResult:
    from docx import Document

    doc = Document(path)
    result = DocxParseResult()

    lines: list[str] = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            lines.append(t)
    result.pages.append(("\n".join(lines), 1))

    for ti, table in enumerate(doc.tables):
        try:
            rows = [[c.text.strip() for c in r.cells] for r in table.rows]
            if len(rows) >= 2:
                result.tables.append(TableBlock(
                    caption=f"（表 {ti + 1}）",
                    headers=rows[0], rows=rows[1:],
                ))
        except Exception:  # noqa: BLE001
            result.warnings.append(f"table {ti + 1}: 解析失败")

    if not lines:
        result.warnings.append("DOCX 无文本内容")
    return result
